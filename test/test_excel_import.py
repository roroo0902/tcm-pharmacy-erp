import datetime as dt
import importlib.util
import unittest
from pathlib import Path
import openpyxl

spec = importlib.util.spec_from_file_location("extract_excel", Path(__file__).parents[1] / "scripts/extract_excel.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def sample(title="測試廠牌 - 測試廠牌", code_column=4):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    sheet.cell(3, code_column, "編號")
    sheet.cell(3, 6, "藥品名稱")
    sheet.cell(3, 7, "價格")
    sheet.cell(3, 8, "備貨量")
    sheet.cell(4, 8, "時間")
    sheet.cell(3, 9, "備貨量")
    sheet.cell(4, 9, "未拆封")
    sheet.cell(3, 10, "備貨量")
    sheet.cell(4, 10, "一週量")
    sheet.cell(5, 1, "單方")
    sheet.cell(6, code_column, "A001 ")
    sheet.cell(6, 6, "測試藥品")
    sheet.cell(6, 8, dt.datetime(2026, 8, 15))
    sheet.cell(6, 9, 10)
    return sheet


class ImportTests(unittest.TestCase):
    def test_dynamic_code_columns(self):
        for col in [2, 4, 5]:
            s = sample(code_column=col)
            records, meta = module.parse_sheet(s, s)
            self.assertEqual(meta["codeColumn"], col)
            self.assertEqual(records[0]["code"], "A001")
            self.assertEqual(records[0]["referenceBottles"], 10)

    def test_date_in_code_not_invented(self):
        s = sample()
        s.cell(6, 4, dt.datetime(2026, 8, 10))
        rows, _ = module.parse_sheet(s, s)
        self.assertEqual(rows[0]["code"], "")
        self.assertEqual(rows[0]["rawCode"], "2026-08-10")
        self.assertIn("缺編號或編號型別錯誤", rows[0]["warnings"])

    def test_zero_text_preserved_and_not_opening(self):
        s = sample()
        s.cell(6, 9, "0無")
        rows, _ = module.parse_sheet(s, s)
        self.assertEqual(rows[0]["rawStock"], "0無")
        self.assertEqual(rows[0]["referenceBottles"], 0)
        self.assertNotIn("openingBottles", rows[0])
        self.assertTrue(rows[0]["warnings"])

    def test_blank_stock_not_zero(self):
        self.assertEqual(module.quantity(None)[0], "")
        self.assertEqual(module.quantity("不明")[0], "")
        self.assertEqual(module.quantity(False)[0], "")

    def test_conflicting_brand_explicit(self):
        s = sample("廠牌甲複方 - 廠牌乙單方")
        rows, meta = module.parse_sheet(s, s)
        self.assertTrue(meta["ambiguousBrand"])
        self.assertIn("分頁廠牌名稱矛盾，待確認", rows[0]["warnings"])

    def test_sections_and_stable_source_key(self):
        s = sample()
        s.cell(7, 1, "複方")
        s.cell(8, 4, "B001")
        s.cell(8, 6, "測試複方")
        s.cell(8, 9, 3)
        rows, _ = module.parse_sheet(s, s)
        again, _ = module.parse_sheet(s, s)
        self.assertEqual([r["category"] for r in rows], ["單方", "複方"])
        self.assertEqual([r["proposedId"] for r in rows], [r["proposedId"] for r in again])
        self.assertEqual(len(set(r["proposedId"] for r in rows)), 2)

    def test_skip_non_drug_sheet(self):
        s = openpyxl.Workbook().active
        s.cell(1, 1, "資料說明")
        rows, meta = module.parse_sheet(s, s)
        self.assertEqual(rows, [])
        self.assertEqual(meta["role"], "非藥品資料頁")


if __name__ == "__main__":
    unittest.main()
