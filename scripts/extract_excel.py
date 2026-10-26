"""Read a Numbers-exported pharmacy XLSX without executing macros or formulas.

Usage: python scripts/extract_excel.py workbook.xlsx --out data/excel-import.json
Requires openpyxl. The input is read-only. Output includes private source data.
"""
import argparse
import datetime as dt
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl


def scalar(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()[:10]
    if value is None:
        return ""
    if isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def quantity(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return value, ""
    if isinstance(value, str) and value.strip() == "0無":
        return 0, "原瓶數為文字0無；僅來源參考值轉為0，需盤點"
    return "", "瓶數空白" if value in (None, "") else "瓶數不是有效數值"


def parse_sheet(sheet, cached):
    header = None
    for row in sheet.iter_rows(min_row=1, max_row=min(8, sheet.max_row), max_col=min(20, sheet.max_column)):
        labels = {str(c.value).strip(): c.column for c in row if c.value is not None}
        if "藥品名稱" in labels and "編號" in labels:
            header = (row[0].row, labels)
            break
    if header is None:
        return [], {"sheet": sheet.title, "role": "非藥品資料頁", "count": 0}
    hr, cols = header
    name_col, code_col = cols["藥品名稱"], cols["編號"]
    stock_cols = [c.column for c in sheet[hr] if c.value == "備貨量"]
    stock_col = next((c for c in stock_cols if sheet.cell(hr + 1, c).value == "未拆封"), None)
    stock_date_col = next((c for c in stock_cols if sheet.cell(hr + 1, c).value == "時間"), None)
    weekly_col = next((c for c in stock_cols if sheet.cell(hr + 1, c).value == "一週量"), None)
    if stock_col is None:
        raise ValueError(f"{sheet.title}: 找不到未拆封瓶數欄")
    prefix, _, suffix = sheet.title.partition(" - ")
    brand = prefix.replace("單方", "").replace("複方", "").strip()
    brand_warning = bool(suffix and suffix not in (prefix, "單方", "複方") and suffix.replace("單方", "").replace("複方", "").strip() != brand)
    category = "單方" if "單方" in prefix else "複方" if "複方" in prefix else ""
    result = []
    for rownum in range(hr + 1, sheet.max_row + 1):
        sections = [sheet.cell(rownum, c).value for c in range(1, name_col)]
        for section in sections:
            if section in ("單方", "複方", "其他"):
                category = section
        raw_name = sheet.cell(rownum, name_col).value
        if raw_name in (None, ""):
            continue
        if not isinstance(raw_name, str) or raw_name.startswith("="):
            raise ValueError(f"{sheet.title}!{rownum}: 藥名不是文字")
        name = raw_name.strip()
        if not name:
            continue
        raw_code = sheet.cell(rownum, code_col).value
        code = raw_code.strip() if isinstance(raw_code, str) else str(raw_code) if isinstance(raw_code, (int, float)) and not isinstance(raw_code, bool) else ""
        if code.startswith("="):
            code = ""
        raw_stock = sheet.cell(rownum, stock_col).value
        qty, qty_note = quantity(cached.cell(rownum, stock_col).value)
        warnings = []
        if not code:
            warnings.append("缺編號或編號型別錯誤")
        if not category:
            warnings.append("類別未辨識")
        if brand_warning:
            warnings.append("分頁廠牌名稱矛盾，待確認")
        if qty_note:
            warnings.append(qty_note)
        if isinstance(qty, (int, float)) and qty < 0:
            warnings.append("來源瓶數為負數")
        if isinstance(raw_stock, str) and raw_stock.startswith("="):
            warnings.append("來源瓶數為公式，參考值使用Excel快取")
        stock_date = scalar(sheet.cell(rownum, stock_date_col).value) if stock_date_col else ""
        if not stock_date:
            warnings.append("來源庫存日期空白")
        source_key = f"{sheet.title}!{rownum}"
        result.append({
            "sourceKey": source_key,
            "proposedId": "X" + hashlib.sha256(source_key.encode()).hexdigest()[:12].upper(),
            "sourceSheet": sheet.title, "sourceRow": rownum,
            "brand": brand, "category": category or "其他", "code": code,
            "rawCode": scalar(raw_code), "name": name, "unit": "瓶",
            "price": scalar(cached.cell(rownum, cols["價格"]).value) if "價格" in cols else "",
            "rawStock": scalar(raw_stock), "referenceBottles": qty,
            "stockDate": stock_date,
            "weeklyUsage": scalar(cached.cell(rownum, weekly_col).value) if weekly_col else "",
            "stockpile": bool(sheet.cell(rownum, cols.get("屯藥", cols.get("囤藥", 1))).value is True),
            "home": bool(sheet.cell(rownum, cols.get("家裡", 1)).value is True),
            "warnings": warnings,
            "sourceColumns": {"code": code_col, "name": name_col, "stock": stock_col},
        })
    return result, {"sheet": sheet.title, "role": "藥品資料頁", "count": len(result), "headerRow": hr, "codeColumn": code_col, "nameColumn": name_col, "stockColumn": stock_col, "ambiguousBrand": brand_warning}


def extract(path):
    book = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    rows, sheets = [], []
    for sheet in book:
        records, meta = parse_sheet(sheet, cached[sheet.title])
        rows.extend(records)
        sheets.append(meta)
    codes = Counter((r["brand"], r["category"], r["code"]) for r in rows if r["code"])
    names = Counter((r["brand"], r["category"], r["name"]) for r in rows)
    for row in rows:
        if row["code"] and codes[(row["brand"], row["category"], row["code"])] > 1:
            row["warnings"].append("同廠牌同類別編號重複，不自動合併")
        if names[(row["brand"], row["category"], row["name"])] > 1:
            row["warnings"].append("同廠牌同類別藥名重複，不自動合併")
    dates = [r["stockDate"] for r in rows if isinstance(r["stockDate"], str) and len(r["stockDate"]) == 10 and r["stockDate"][4:5] == "-"]
    return {"sourceFile": Path(path).name, "sourceSha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(), "schemaVersion": 1,
            "summary": {"sheetCount": len(sheets), "recordCount": len(rows), "missingCodeCount": sum(not r["code"] for r in rows), "warningRecordCount": sum(bool(r["warnings"]) for r in rows), "stockDateMin": min(dates, default=""), "stockDateMax": max(dates, default="")},
            "sheets": sheets, "records": rows}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    result = extract(args.source)
    dest = Path(args.out)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"summary": result["summary"], "sheets": result["sheets"]}, ensure_ascii=False))
